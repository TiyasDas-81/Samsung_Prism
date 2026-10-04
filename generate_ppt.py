import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# Helper to add a colored text box
def add_textbox(slide, left, top, width, height, text, font_size, font_bold, font_color, align=PP_ALIGN.LEFT, font_name="Arial"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = font_bold
    run.font.color.rgb = font_color
    run.font.name = font_name
    return txBox

def main():
    prs = Presentation()
    
    # Set slide dimensions (widescreen 16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    bg_color = RGBColor(10, 15, 30) # Very dark navy/near-black
    text_light = RGBColor(240, 240, 240)
    accent_cyan = RGBColor(0, 255, 255)
    accent_blue = RGBColor(0, 102, 255)
    accent_warning = RGBColor(255, 102, 0)
    card_bg = RGBColor(20, 28, 50)
    
    # Helper to apply background
    def apply_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
    blank_slide_layout = prs.slide_layouts[6]
    
    # --------------------------------------------------
    # SLIDE 1 - HERO
    # --------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide1)
    
    add_textbox(slide1, Inches(1), Inches(2), Inches(11.33), Inches(1.5), 
                "Interruptible Real-Time Agents", Pt(54), True, text_light, PP_ALIGN.CENTER)
                
    add_textbox(slide1, Inches(1), Inches(3.5), Inches(11.33), Inches(1), 
                "Full-Duplex Voice Coordination with Safe Interruption Recovery", Pt(28), False, accent_cyan, PP_ALIGN.CENTER)
                
    # Central glowing node concept (using shapes)
    shape = slide1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), Inches(4.5), Inches(1.33), Inches(1.33))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_blue
    shape.line.width = Pt(2)
    
    # Team Block
    shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(4.9), Inches(3.5), Inches(1.6))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(15, 20, 35)
    shape.line.color.rgb = accent_blue
    shape.line.width = Pt(1)
    
    add_textbox(slide1, Inches(1.1), Inches(5.0), Inches(3.3), Inches(0.3), "KKR - Kolkata Kode Riders", Pt(14), True, accent_cyan, font_name="Segoe UI")
    add_textbox(slide1, Inches(1.1), Inches(5.3), Inches(3.3), Inches(0.3), "Arhit Basu — Leader", Pt(14), True, text_light, font_name="Segoe UI")
    add_textbox(slide1, Inches(1.1), Inches(5.6), Inches(3.3), Inches(0.8), "Tiyas Das\nSubarta Ghosh\nSoumen Mondal", Pt(13), False, text_light, font_name="Segoe UI")
    
    add_textbox(slide1, Inches(1), Inches(6.8), Inches(11.33), Inches(0.5), 
                "Samsung PRISM GenAI Hackathon 2026 • Theme 05\nGitHub: TiyasDas-81/Samsung_Prism", Pt(12), False, text_light, PP_ALIGN.CENTER)

    # --------------------------------------------------
    # SLIDE 2 - THE PROBLEM
    # --------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide2)
    add_textbox(slide2, Inches(1), Inches(0.5), Inches(11.33), Inches(1), 
                "Voice Agents Break When Users Change Their Mind", Pt(40), True, text_light)
                
    # Flow
    add_textbox(slide2, Inches(1), Inches(1.8), Inches(4), Inches(0.5), "USER", Pt(20), True, accent_cyan)
    add_textbox(slide2, Inches(1), Inches(2.2), Inches(4), Inches(0.5), "\"Navigate to Chennai Central...\"", Pt(18), False, text_light)
    
    add_textbox(slide2, Inches(1), Inches(2.8), Inches(4), Inches(0.5), "↓", Pt(20), True, accent_cyan)
    
    add_textbox(slide2, Inches(1), Inches(3.3), Inches(4), Inches(0.5), "AGENT", Pt(20), True, accent_blue)
    add_textbox(slide2, Inches(1), Inches(3.7), Inches(4), Inches(0.5), "\"Sure, navigating...\"", Pt(18), False, text_light)
    
    add_textbox(slide2, Inches(1), Inches(4.3), Inches(4), Inches(0.5), "↓", Pt(20), True, accent_cyan)
    
    add_textbox(slide2, Inches(1), Inches(4.8), Inches(5), Inches(0.5), "USER INTERRUPTS", Pt(20), True, accent_warning)
    add_textbox(slide2, Inches(1), Inches(5.2), Inches(5), Inches(0.5), "\"...wait, I meant VIT Vellore.\"", Pt(18), False, text_light)
    
    # Naive agent box
    shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(2), Inches(4.5), Inches(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_warning
    
    add_textbox(slide2, Inches(7.7), Inches(2.2), Inches(4.1), Inches(1.5),
                "NAIVE AGENT\n❌ stale action continues\n❌ wrong tool state\n❌ poor UX", Pt(18), False, text_light)
                
    # Core problem statement
    add_textbox(slide2, Inches(7.5), Inches(4.5), Inches(4.5), Inches(1.5),
                "THE CORE PROBLEM:\nOld work must not survive a new intent.", Pt(24), True, accent_warning)

    # --------------------------------------------------
    # SLIDE 3 - OUR SOLUTION
    # --------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide3)
    add_textbox(slide3, Inches(1), Inches(0.5), Inches(11.33), Inches(1), 
                "Make Interruption a First-Class Event", Pt(40), True, text_light)
                
    cards = [
        ("01 Coordination Layer", "Owns the conversation state"),
        ("02 Fast / Slow Path", "Respond immediately, execute asynchronously"),
        ("03 Interrupt Detection", "Cancel obsolete work"),
        ("04 State Generations", "Reject stale results"),
        ("05 Idempotent Tools", "Prevent duplicate mutations")
    ]
    
    for i, (title, desc) in enumerate(cards):
        y = 1.8 + (i * 0.9)
        shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(y), Inches(4.5), Inches(0.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_bg
        shape.line.color.rgb = accent_blue
        add_textbox(slide3, Inches(1.1), Inches(y+0.1), Inches(4.3), Inches(0.3), title, Pt(18), True, accent_cyan)
        add_textbox(slide3, Inches(1.1), Inches(y+0.4), Inches(4.3), Inches(0.3), desc, Pt(14), False, text_light)
        
    shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2.5), Inches(5.5), Inches(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_cyan
    shape.line.width = Pt(2)
    add_textbox(slide3, Inches(6.5), Inches(3.2), Inches(5.5), Inches(0.5), "INTERRUPTION-AWARE COORDINATOR", Pt(20), True, text_light, PP_ALIGN.CENTER)
    
    add_textbox(slide3, Inches(1), Inches(6.5), Inches(11.33), Inches(0.5), 
                "NEW INTENT → CANCEL OLD WORK → INCREMENT GENERATION → ACCEPT ONLY FRESH RESULTS", Pt(16), True, accent_cyan, PP_ALIGN.CENTER)

    # --------------------------------------------------
    # SLIDE 4 - SYSTEM ARCHITECTURE (REDESIGNED)
    # --------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide4)
    add_textbox(slide4, Inches(1), Inches(0.3), Inches(11.33), Inches(1), 
                "Architecture: From Speech to Safe Tool Execution", Pt(40), True, text_light)
                
    # Vertical architecture spacing
    y_starts = [1.2, 1.75, 2.3, 2.85, 4.65, 5.2, 5.75, 6.3]
    heights = [0.4, 0.4, 0.4, 1.75, 0.4, 0.4, 0.4, 0.4]
    widths = [3.5, 3.5, 3.5, 4.5, 3.5, 3.5, 3.5, 3.5]
    xs = [4.91, 4.91, 4.91, 4.41, 4.91, 4.91, 4.91, 4.91]
    
    nodes = [
        "USER VOICE", 
        "LIVEKIT VAD", 
        "STT", 
        "INTERRUPTION-AWARE COORDINATION LAYER\ngeneration_id\ncancellation\nstate validation", 
        "LLM", 
        "ASYNC TOOL EXECUTION", 
        "FRESH RESULT ONLY", 
        "TTS"
    ]
    
    for i in range(8):
        shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(xs[i]), Inches(y_starts[i]), Inches(widths[i]), Inches(heights[i]))
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_bg
        
        if i == 3: # Coordination layer
            shape.line.color.rgb = accent_cyan
            add_textbox(slide4, Inches(xs[i]+0.1), Inches(y_starts[i]+0.05), Inches(widths[i]-0.2), Inches(1.5), nodes[i], Pt(13), True, text_light, PP_ALIGN.CENTER)
        else:
            shape.line.color.rgb = accent_blue
            add_textbox(slide4, Inches(xs[i]+0.1), Inches(y_starts[i]+0.05), Inches(widths[i]-0.2), Inches(0.3), nodes[i], Pt(14), True, text_light, PP_ALIGN.CENTER)

        # Arrows (downward)
        if i < 7:
            arrow_y = y_starts[i] + heights[i]
            arrow_h = y_starts[i+1] - arrow_y
            add_textbox(slide4, Inches(6.5), Inches(arrow_y-0.1), Inches(0.33), Inches(arrow_h+0.2), "↓", Pt(14), True, text_light, PP_ALIGN.CENTER)

    # Interruption branch (Left)
    shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.95), Inches(2.6), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_warning
    add_textbox(slide4, Inches(1.6), Inches(3.05), Inches(2.4), Inches(1.0), "INTERRUPTION\n↓\nCANCEL PENDING TASK\n↓\ngeneration++", Pt(13), True, accent_warning, PP_ALIGN.CENTER)

    # Stale branch (Right)
    shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(5.4), Inches(2.6), Inches(0.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_warning
    add_textbox(slide4, Inches(9.3), Inches(5.45), Inches(2.4), Inches(0.7), "STALE RESULT\n↓\nREJECT", Pt(13), True, accent_warning, PP_ALIGN.CENTER)

    # --------------------------------------------------
    # SLIDE 5 - FDB-v3 + RESULTS
    # --------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide5)
    add_textbox(slide5, Inches(1), Inches(0.5), Inches(11.33), Inches(1), 
                "Evaluated Against Full-Duplex-Bench v3", Pt(40), True, text_light)
                
    add_textbox(slide5, Inches(1), Inches(1.8), Inches(11.33), Inches(0.5), 
                "100 scenarios → real human disfluencies → LiveKit agent → tool calls → evaluator", Pt(18), True, accent_cyan, PP_ALIGN.CENTER)
                
    # Metric cards
    metrics = [("77.3%", "Tool Selection"), ("47.0%", "Argument Accuracy"), ("33.3%", "Strict Pass Rate")]
    for i, (val, lbl) in enumerate(metrics):
        x = 1.5 + (i * 3.6)
        shape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.8), Inches(3.2), Inches(2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_bg
        shape.line.color.rgb = accent_blue
        
        add_textbox(slide5, Inches(x), Inches(3.1), Inches(3.2), Inches(1), val, Pt(48), True, accent_cyan, PP_ALIGN.CENTER)
        add_textbox(slide5, Inches(x), Inches(4.1), Inches(3.2), Inches(0.5), lbl, Pt(18), False, text_light, PP_ALIGN.CENTER)
        
    shape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(5.3), Inches(5.33), Inches(0.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(50, 40, 10)
    shape.line.color.rgb = accent_warning
    
    add_textbox(slide5, Inches(4), Inches(5.4), Inches(5.33), Inches(0.6), "LOCAL / PARTIAL RUN\n30 / 100 scenarios completed", Pt(20), True, accent_warning, PP_ALIGN.CENTER)
    
    add_textbox(slide5, Inches(1), Inches(6.5), Inches(11.33), Inches(0.5), "Run terminated at scenario #30 due to Groq API 429 rate limiting.", Pt(14), False, text_light, PP_ALIGN.CENTER)

    # --------------------------------------------------
    # SLIDE 6 - EXTENSION USE CASE
    # --------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide6)
    add_textbox(slide6, Inches(1), Inches(0.5), Inches(11.33), Inches(1), 
                "Beyond the Benchmark: In-Car Destination Change", Pt(40), True, text_light)
                
    # Left column flow
    add_textbox(slide6, Inches(1), Inches(1.8), Inches(4), Inches(0.5), "📍 CHENNAI CENTRAL", Pt(20), True, accent_cyan)
    add_textbox(slide6, Inches(1), Inches(2.3), Inches(4), Inches(0.5), "↓\n\"Navigating...\"", Pt(16), False, text_light)
    add_textbox(slide6, Inches(1), Inches(3.1), Inches(4), Inches(0.5), "↓\n🔊 USER INTERRUPTS", Pt(16), True, accent_warning)
    add_textbox(slide6, Inches(1), Inches(3.9), Inches(4), Inches(0.5), "\"Wait — VIT Vellore.\"", Pt(16), False, text_light)
    add_textbox(slide6, Inches(1), Inches(4.5), Inches(4), Inches(0.5), "↓\n✕ STALE ROUTE", Pt(16), True, accent_warning)
    add_textbox(slide6, Inches(1), Inches(5.3), Inches(4), Inches(0.5), "↓\n✓ NEW INTENT\n\"VIT Vellore\"", Pt(16), True, accent_cyan)
    add_textbox(slide6, Inches(1), Inches(6.2), Inches(4), Inches(0.5), "↓\n📍 VIT VELLORE", Pt(20), True, accent_cyan)

    # Right column graphic
    shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(2.3), Inches(6), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_warning
    add_textbox(slide6, Inches(6.2), Inches(2.5), Inches(5.6), Inches(0.8), "OLD GENERATION\n──────────────X discarded", Pt(20), True, accent_warning)

    shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(4.0), Inches(6), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.color.rgb = accent_cyan
    add_textbox(slide6, Inches(6.2), Inches(4.2), Inches(5.6), Inches(0.8), "NEW GENERATION\n──────────────✓ executed", Pt(20), True, accent_cyan)
    
    add_textbox(slide6, Inches(6), Inches(5.8), Inches(6), Inches(1), "New intent wins.\nOld work cannot mutate current state.", Pt(24), True, text_light)

    # --------------------------------------------------
    # SLIDE 7 - DEMO / REPRODUCIBILITY
    # --------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide7)
    add_textbox(slide7, Inches(1), Inches(0.5), Inches(11.33), Inches(1), 
                "Built to Run, Interrupt, Recover", Pt(40), True, text_light)
                
    cols = [
        ("RUN", ["LiveKit", "Groq", "ElevenLabs"]),
        ("INTERRUPT", ["VAD trigger", "cancellation", "generation update"]),
        ("RECOVER", ["stale-result rejection", "validated tool call", "corrected response"])
    ]
    
    for i, (title, items) in enumerate(cols):
        x = 1 + (i * 3.9)
        shape = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.3), Inches(3.5), Inches(2.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_bg
        shape.line.color.rgb = accent_blue
        
        add_textbox(slide7, Inches(x), Inches(2.5), Inches(3.5), Inches(0.5), title, Pt(24), True, accent_cyan, PP_ALIGN.CENTER)
        
        y = 3.3
        for item in items:
            add_textbox(slide7, Inches(x+0.2), Inches(y), Inches(3.1), Inches(0.4), "• " + item, Pt(18), False, text_light)
            y += 0.5
            
    add_textbox(slide7, Inches(1), Inches(5.3), Inches(11.33), Inches(0.5), "Python 3.11 • LiveKit Agents 1.8.3", Pt(18), True, text_light, PP_ALIGN.CENTER)
    
    shape = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(6.1), Inches(9.33), Inches(0.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 30, 30)
    shape.line.color.rgb = text_light
    add_textbox(slide7, Inches(2), Inches(6.25), Inches(9.33), Inches(0.5), "GPT-4o-dependent latency / LLM-judge metrics were unavailable locally and are not claimed.", Pt(14), False, text_light, PP_ALIGN.CENTER)

    # --------------------------------------------------
    # SLIDE 8 - CLOSING
    # --------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide8)
    add_textbox(slide8, Inches(1), Inches(0.5), Inches(11.33), Inches(1.2), 
                "Interruption Should Change the Action —\nNot Break the Agent.", Pt(36), True, text_light)
                
    # Center graphic
    add_textbox(slide8, Inches(4), Inches(2.0), Inches(5.33), Inches(2), "VOICE\n↓\nINTENT\n↓\nCOORDINATION\n↓\nSAFE TOOL EXECUTION", Pt(18), True, accent_cyan, PP_ALIGN.CENTER)
    
    # Takeaways
    add_textbox(slide8, Inches(1), Inches(4.3), Inches(11.33), Inches(1.2), 
                "✓ Interruptions are first-class events\n✓ Stale work is safely rejected\n✓ The architecture extends beyond benchmark scenarios", Pt(20), False, text_light, PP_ALIGN.CENTER)
                
    # Team Footer
    add_textbox(slide8, Inches(1), Inches(5.8), Inches(11.33), Inches(0.4), 
                "KKR - Kolkata Kode Riders: Arhit Basu (Leader) • Tiyas Das • Subarta Ghosh • Soumen Mondal", Pt(14), False, text_light, PP_ALIGN.CENTER, font_name="Segoe UI")
    
    # AI Disclosure / Repos
    add_textbox(slide8, Inches(1), Inches(6.4), Inches(11.33), Inches(0.8), 
                "Samsung PRISM GenAI Hackathon 2026 • Theme 05\nGitHub: TiyasDas-81/Samsung_Prism\n\nDeveloped with assistance from Google DeepMind Antigravity AI for architecture design and testing.", Pt(12), False, text_light, PP_ALIGN.CENTER)

    # Save
    import os
    if not os.path.exists("presentation"):
        os.makedirs("presentation")
    prs.save("presentation/Samsung_PRISM_Theme05.pptx")
    print("Presentation created successfully.")

if __name__ == "__main__":
    main()
